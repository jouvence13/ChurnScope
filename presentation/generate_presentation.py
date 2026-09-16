#!/usr/bin/env python3
"""
Generateur du support de soutenance ChurnScope (15 diapositives, format 16:9).

Ce script reconstruit integralement le fichier soutenance_churnscope.pptx a partir
des resultats reels et verifies du projet (voir docs/01_analyse_preparation.md,
docs/02_apprentissage_modeles.md, docs/03_validation_conclusion.md et le README).
Il ne depend d'aucun fichier binaire existant : relancer ce script recree la
presentation a l'identique, ce qui la rend reproductible et versionnable en toute
securite (le .pptx final reste volontairement hors du depot Git, voir .gitignore).

Prerequis : pip install python-pptx

Utilisation :
    python presentation/generate_presentation.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --------------------------------------------------------------------------- #
# Palette et constantes de mise en page
# --------------------------------------------------------------------------- #

NAVY = RGBColor(0x1E, 0x3A, 0x5F)  # Bleu principal (titres, blocs forts)
NAVY_DARK = RGBColor(0x15, 0x28, 0x42)  # Fond des diapositives sombres
GOLD = RGBColor(0xC9, 0x8A, 0x2C)  # Accent (mise en avant, modele retenu)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF2, 0xF4, 0xF7)
DARK_TEXT = RGBColor(0x22, 0x2A, 0x35)
MID_GRAY = RGBColor(0x5B, 0x64, 0x70)
ICE = RGBColor(0xB9, 0xC7, 0xD8)  # Texte secondaire sur fond sombre
GREEN_TEXT = RGBColor(0x1F, 0x6E, 0x49)
GREEN_BG = RGBColor(0xE3, 0xF1, 0xE7)
CORAL_TEXT = RGBColor(0xA6, 0x36, 0x2A)
CORAL_BG = RGBColor(0xFB, 0xE8, 0xE5)

SLIDE_W = Inches(10)
SLIDE_H = Inches(5.63)

PRESENTATION_DIR = Path(__file__).resolve().parent
OUTPUT_PATH = PRESENTATION_DIR / "soutenance_churnscope.pptx"


# --------------------------------------------------------------------------- #
# Fonctions utilitaires de mise en forme
# --------------------------------------------------------------------------- #


def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _blank(prs, bg_color):
    layout = prs.slide_layouts[6]  # layout vierge
    slide = prs.slides.add_slide(layout)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = bg_color
    bg.line.fill.background()
    bg.shadow.inherit = False
    spTree = slide.shapes._spTree
    spTree.remove(bg._element)
    spTree.insert(2, bg._element)
    return slide


def light_slide(prs):
    return _blank(prs, WHITE)


def dark_slide(prs):
    return _blank(prs, NAVY_DARK)


def add_textbox(
    slide,
    left,
    top,
    width,
    height,
    text,
    *,
    size=14,
    bold=False,
    italic=False,
    color=DARK_TEXT,
    align=PP_ALIGN.LEFT,
    font="Calibri",
    anchor=None,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    if anchor is not None:
        tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    return box


def set_shape_text(
    shape,
    text,
    *,
    size=14,
    bold=False,
    color=DARK_TEXT,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.MIDDLE,
):
    """Ecrit un texte dans une forme deja existante (au lieu d'une nouvelle zone de texte)."""
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return shape


def add_bullets(
    slide, left, top, width, height, items, *, size=13, color=DARK_TEXT, bullet_char="—"
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(7)
        run = p.add_run()
        run.text = f"{bullet_char}  {item}"
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = "Calibri"
    return box


def add_pill(slide, left, top, width, height, text, *, fill=NAVY, color=WHITE, size=11):
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    pill.adjustments[0] = 0.5
    pill.fill.solid()
    pill.fill.fore_color.rgb = fill
    pill.line.fill.background()
    pill.shadow.inherit = False
    set_shape_text(pill, text, size=size, bold=True, color=color, align=PP_ALIGN.CENTER)
    return pill


def add_header(slide, kicker, title, *, dark=False):
    """Amorce commune de titre pour les diapositives de contenu (sans bandeau plein largeur)."""
    kicker_color = GOLD if not dark else GOLD
    title_color = NAVY if not dark else WHITE
    add_textbox(
        slide,
        Inches(0.5),
        Inches(0.42),
        Inches(9),
        Inches(0.3),
        kicker.upper(),
        size=12,
        bold=True,
        color=kicker_color,
        font="Calibri",
    )
    add_textbox(
        slide,
        Inches(0.5),
        Inches(0.72),
        Inches(9),
        Inches(0.6),
        title,
        size=28,
        bold=True,
        color=title_color,
    )


def add_stat_tile(
    slide,
    left,
    top,
    width,
    height,
    value,
    label,
    *,
    accent=NAVY,
    fill=LIGHT_GRAY,
    value_size=19,
):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    box.adjustments[0] = 0.09
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.color.rgb = accent
    box.line.width = Pt(1)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    p1 = tf.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    r1 = p1.add_run()
    r1.text = value
    r1.font.size = Pt(value_size)
    r1.font.bold = True
    r1.font.color.rgb = accent
    r1.font.name = "Calibri"
    if label:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = label
        r2.font.size = Pt(10)
        r2.font.color.rgb = MID_GRAY
        r2.font.name = "Calibri"


def add_quote_box(
    slide, left, top, width, height, text, *, fill=LIGHT_GRAY, color=DARK_TEXT, size=13
):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    box.adjustments[0] = 0.12
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.fill.background()
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.25)
    tf.margin_right = Inches(0.25)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.italic = True
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return box


def add_table(
    slide,
    left,
    top,
    width,
    height,
    headers,
    rows,
    *,
    highlight_rows=(),
    font_size=12,
    first_col_left=True,
    col_widths=None,
):
    n_rows = len(rows) + 1
    n_cols = len(headers)
    table_shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    table = table_shape.table

    if col_widths is not None:
        for c, w in enumerate(col_widths):
            table.columns[c].width = w

    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        cell.margin_top = Pt(4)
        cell.margin_bottom = Pt(4)
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if (c == 0 and first_col_left) else PP_ALIGN.CENTER
        run = p.add_run()
        run.text = h
        run.font.bold = True
        run.font.color.rgb = WHITE
        run.font.size = Pt(font_size)

    for r, row in enumerate(rows, start=1):
        highlighted = (r - 1) in highlight_rows
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = (
                RGBColor(0xFB, 0xF2, 0xDF) if highlighted else WHITE
            )
            cell.margin_top = Pt(4)
            cell.margin_bottom = Pt(4)
            p = cell.text_frame.paragraphs[0]
            p.alignment = (
                PP_ALIGN.LEFT if (c == 0 and first_col_left) else PP_ALIGN.CENTER
            )
            p.runs[0].font.size = Pt(font_size)
            p.runs[0].font.bold = highlighted
            p.runs[0].font.color.rgb = GOLD if highlighted else DARK_TEXT
    return table


def add_footer(slide, page_no, *, dark=False):
    color = ICE if dark else MID_GRAY
    add_textbox(
        slide,
        Inches(0.5),
        SLIDE_H - Inches(0.32),
        Inches(6),
        Inches(0.3),
        "ChurnScope — Prediction et prevention du churn client",
        size=9,
        color=color,
    )
    add_textbox(
        slide,
        SLIDE_W - Inches(1.3),
        SLIDE_H - Inches(0.32),
        Inches(0.9),
        Inches(0.3),
        str(page_no),
        size=9,
        color=color,
        align=PP_ALIGN.RIGHT,
    )


def add_flow_chips(
    slide, top, items, *, chip_w=Inches(1.98), chip_h=Inches(1.0), gap=Inches(0.32)
):
    """Chaine horizontale de blocs relies par des fleches (ex. cadrage, parcours)."""
    x = Inches(0.5)
    for i, (line1, line2) in enumerate(items):
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, x, top, chip_w, chip_h
        )
        box.adjustments[0] = 0.08
        box.fill.solid()
        box.fill.fore_color.rgb = NAVY
        box.line.fill.background()
        box.shadow.inherit = False
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.1)
        tf.margin_right = Inches(0.1)
        p1 = tf.paragraphs[0]
        p1.alignment = PP_ALIGN.CENTER
        r1 = p1.add_run()
        r1.text = line1
        r1.font.size = Pt(12)
        r1.font.bold = True
        r1.font.color.rgb = WHITE
        r1.font.name = "Calibri"
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = line2
        r2.font.size = Pt(9)
        r2.font.color.rgb = ICE
        r2.font.name = "Calibri"
        x = Emu(int(x) + int(chip_w) + int(gap))
        if i < len(items) - 1:
            add_textbox(
                slide,
                Emu(int(x) - int(gap)),
                Emu(int(top) + int(chip_h) // 2 - int(Inches(0.2))),
                gap,
                Inches(0.4),
                "→",
                size=16,
                bold=True,
                color=GOLD,
                align=PP_ALIGN.CENTER,
            )


def add_step_list(
    slide, left, top, width, items, *, step_h=Inches(0.44), gap=Inches(0.13), size=11
):
    """Liste verticale d'etapes numerotees (bandes alternees)."""
    y = top
    for i, step in enumerate(items):
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, y, width, step_h
        )
        box.adjustments[0] = 0.18
        box.fill.solid()
        box.fill.fore_color.rgb = NAVY if i % 2 == 0 else GOLD
        box.line.fill.background()
        box.shadow.inherit = False
        set_shape_text(
            box,
            f"{i + 1}.  {step}",
            size=size,
            bold=True,
            color=WHITE,
            align=PP_ALIGN.LEFT,
        )
        box.text_frame.margin_left = Inches(0.2)
        box.text_frame.margin_right = Inches(0.15)
        y = Emu(int(y) + int(step_h) + int(gap))
    return y


# --------------------------------------------------------------------------- #
# Diapositives
# --------------------------------------------------------------------------- #


def slide_01_title(prs):
    slide = dark_slide(prs)

    add_pill(
        slide,
        Inches(0.75),
        Inches(0.7),
        Inches(3.75),
        Inches(0.42),
        "SOUTENANCE AIA01 — SEPTEMBRE 2026",
        fill=GOLD,
        color=NAVY_DARK,
        size=10.5,
    )

    add_textbox(
        slide,
        Inches(0.75),
        Inches(1.5),
        Inches(8.5),
        Inches(0.85),
        "ChurnScope",
        size=44,
        bold=True,
        color=WHITE,
    )
    add_textbox(
        slide,
        Inches(0.78),
        Inches(2.35),
        Inches(8.5),
        Inches(0.5),
        "Prediction et prevention du depart client",
        size=19,
        color=ICE,
    )
    add_textbox(
        slide,
        Inches(0.78),
        Inches(2.9),
        Inches(8.5),
        Inches(0.5),
        "« Qui devons-nous contacter cette semaine, et pourquoi ? »",
        size=14,
        italic=True,
        color=ICE,
    )

    tags = [
        "TELECOMMUNICATIONS",
        "TELCO CUSTOMER CHURN · KAGGLE",
        "RANDOM FOREST OPTIMISE",
    ]
    x = Inches(0.78)
    for tag in tags:
        w = Inches(0.095 * len(tag) + 0.55)
        add_pill(
            slide, x, Inches(3.6), w, Inches(0.36), tag, fill=NAVY, color=ICE, size=9.5
        )
        x = Emu(int(x) + int(w) + int(Inches(0.15)))

    add_textbox(
        slide,
        Inches(0.78),
        Inches(4.55),
        Inches(6),
        Inches(0.35),
        "Jouvence",
        size=16,
        bold=True,
        color=WHITE,
    )
    add_textbox(
        slide,
        Inches(0.78),
        Inches(4.9),
        Inches(7),
        Inches(0.3),
        "Directeur de Projet en Intelligence Artificielle — Annee 1 · L'Ecole Multimedia",
        size=11,
        color=ICE,
    )
    add_textbox(
        slide,
        SLIDE_W - Inches(3.3),
        Inches(4.9),
        Inches(2.8),
        Inches(0.3),
        "github.com/jouvence13/ChurnScope",
        size=10,
        color=ICE,
        align=PP_ALIGN.RIGHT,
    )
    return slide


def slide_02_cadrage(prs):
    slide = light_slide(prs)
    add_header(slide, "Cadrage", "Prioriser les clients a contacter")

    add_flow_chips(
        slide,
        Inches(1.35),
        [
            ("Donnees client", "Contrat, services, anciennete, facturation"),
            ("Modele", "Random Forest — score de risque"),
            ("Score & seuil", "Liste de clients au-dessus du seuil"),
            ("Examen humain", "La decision de contact reste a l'equipe"),
        ],
    )

    add_bullets(
        slide,
        Inches(0.5),
        Inches(2.65),
        Inches(9),
        Inches(1.4),
        [
            "Probleme : identifier les profils associes a un risque de resiliation.",
            "Sortie : un score de risque et une liste de clients a examiner selon un seuil.",
            "Compromis : detecter un maximum de departs sans multiplier les fausses alertes.",
        ],
        size=13.5,
    )

    add_quote_box(
        slide,
        Inches(0.5),
        Inches(4.15),
        Inches(9),
        Inches(0.95),
        "Les donnees decrivent des departs deja observes : le projet n'a mesure ni "
        "une reduction reelle du churn en production, ni de retour sur investissement.",
    )
    add_footer(slide, 2)
    return slide


def slide_03_data(prs):
    slide = light_slide(prs)
    add_header(slide, "Donnees", "Un desequilibre a ne pas ignorer")

    stats = [
        ("7 043", "clients, 21 colonnes"),
        ("26,54 %", "taux de churn"),
        ("1 869", "clients partis"),
        ("5 174", "clients restes"),
    ]
    box_w = Inches(2.1)
    gap = Inches(0.15)
    x = Inches(0.5)
    for value, label in stats:
        add_stat_tile(slide, x, Inches(1.35), box_w, Inches(1.05), value, label)
        x = Emu(int(x) + int(box_w) + int(gap))

    add_textbox(
        slide,
        Inches(0.5),
        Inches(2.7),
        Inches(9),
        Inches(0.3),
        "Quatre dimensions de variables",
        size=13,
        bold=True,
        color=NAVY,
    )
    add_bullets(
        slide,
        Inches(0.5),
        Inches(3.05),
        Inches(4.4),
        Inches(1.1),
        [
            "Demographiques : genre, senior, partenaire",
            "Services : telephonie, internet, streaming",
        ],
        size=12.5,
    )
    add_bullets(
        slide,
        Inches(5.1),
        Inches(3.05),
        Inches(4.4),
        Inches(1.1),
        [
            "Compte client : anciennete, contrat, facturation",
            "Cible : Churn (Yes / No)",
        ],
        size=12.5,
    )

    add_quote_box(
        slide,
        Inches(0.5),
        Inches(4.15),
        Inches(9),
        Inches(0.95),
        "Predire systematiquement « pas de churn » donnerait 73,46 % d'accuracy "
        "sans detecter aucun depart — d'ou l'attention portee au rappel, a la "
        "precision, au F1-score et a l'AUC.",
    )
    add_footer(slide, 3)
    return slide


def slide_04_pipeline(prs):
    slide = light_slide(prs)
    add_header(slide, "Methode", "Preparer les donnees sans fuite d'information")

    steps = [
        "Nettoyage — suppression de customerID, conversion de TotalCharges",
        "Valeurs manquantes — imputation par la mediane, apres le split",
        "Encodage & normalisation — One-Hot, StandardScaler",
        "Decoupage stratifie — 80 % train / 20 % test (stratify=y)",
        "Pipeline scikit-learn — reappris a chaque pli de validation croisee",
    ]
    add_step_list(slide, Inches(0.5), Inches(1.3), Inches(9), steps)

    add_quote_box(
        slide,
        Inches(0.5),
        Inches(4.35),
        Inches(9),
        Inches(0.75),
        "« Le jeu de test n'intervient ni dans l'apprentissage des transformations, "
        "ni dans le choix des hyperparametres. »",
    )
    add_footer(slide, 4)
    return slide


def slide_05_protocol(prs):
    slide = light_slide(prs)
    add_header(slide, "Protocole", "Comparer quatre modeles avec la meme rigueur")

    add_table(
        slide,
        Inches(0.5),
        Inches(1.3),
        Inches(9),
        Inches(1.9),
        ["Modele", "Role dans la comparaison"],
        [
            ["Regression logistique", "Reference simple, coefficients interpretables"],
            ["Ridge Classifier", "Comparaison supplementaire (regularisation L2)"],
            ["Arbre de decision", "Regles explicites, relations non lineaires simples"],
            ["Random Forest ★", "Combinaison d'arbres — modele retenu"],
        ],
        highlight_rows={3},
        font_size=12.5,
    )

    add_bullets(
        slide,
        Inches(0.5),
        Inches(3.5),
        Inches(9),
        Inches(1.2),
        [
            "Recherche d'hyperparametres par GridSearchCV, validation croisee stratifiee a 5 plis, "
            "critere : le rappel (recall).",
            "Le jeu de test ne sert qu'a l'evaluation finale, jamais au choix des hyperparametres.",
        ],
        size=13,
    )
    add_footer(slide, 5)
    return slide


def slide_06_comparison(prs):
    slide = light_slide(prs)
    add_header(slide, "Resultats", "Comparer les quatre modeles")

    add_table(
        slide,
        Inches(0.5),
        Inches(1.3),
        Inches(9),
        Inches(2.05),
        ["Modele", "Accuracy", "Precision", "Rappel", "F1-score", "AUC"],
        [
            ["Regression logistique", "0,738", "0,504", "0,783", "0,614", "0,841"],
            ["Ridge Classifier", "0,737", "0,503", "0,789", "0,615", "0,836"],
            ["Arbre de decision", "0,755", "0,527", "0,759", "0,622", "0,832"],
            ["Random Forest ★", "0,755", "0,526", "0,791", "0,632", "0,842"],
        ],
        highlight_rows={3},
        font_size=11.5,
        col_widths=[
            Inches(2.3),
            Inches(1.34),
            Inches(1.34),
            Inches(1.34),
            Inches(1.34),
            Inches(1.34),
        ],
    )

    add_textbox(
        slide,
        Inches(0.5),
        Inches(3.65),
        Inches(9),
        Inches(0.3),
        "Pourquoi le Random Forest ?",
        size=13,
        bold=True,
        color=NAVY,
    )
    add_textbox(
        slide,
        Inches(0.5),
        Inches(4.0),
        Inches(9),
        Inches(0.7),
        "Meilleur rappel (0,791), meilleur F1-score (0,632) et meilleure AUC (0,842) : "
        "c'est le modele retenu pour l'optimisation.",
        size=12.5,
        color=DARK_TEXT,
    )
    add_footer(slide, 6)
    return slide


def slide_07_optimization(prs):
    slide = light_slide(prs)
    add_header(slide, "Optimisation", "Random Forest optimise par GridSearchCV")

    params = [
        "n_estimators = 200",
        "max_depth = 6",
        "min_samples_leaf = 1",
        'class_weight = "balanced"',
    ]
    x = Inches(0.5)
    box_w = Inches(2.15)
    for param in params:
        add_stat_tile(
            slide,
            x,
            Inches(1.3),
            box_w,
            Inches(0.85),
            param,
            "",
            accent=NAVY,
            value_size=13,
        )
        x = Emu(int(x) + int(box_w) + int(Inches(0.1)))

    add_textbox(
        slide,
        Inches(0.5),
        Inches(2.45),
        Inches(6),
        Inches(0.3),
        "Methode",
        size=13,
        bold=True,
        color=NAVY,
    )
    add_bullets(
        slide,
        Inches(0.5),
        Inches(2.8),
        Inches(8.8),
        Inches(1.15),
        [
            "Validation croisee stratifiee a 5 plis, recherche par grille (GridSearchCV)",
            "Optimisation du rappel (recall) : rappel moyen en validation croisee = 0,809",
        ],
        size=13,
    )

    add_quote_box(
        slide,
        Inches(0.5),
        Inches(4.15),
        Inches(9),
        Inches(0.9),
        "Un churn non detecte peut representer une perte commerciale plus importante "
        "qu'une fausse alerte — d'ou la priorite donnee au rappel.",
    )
    add_footer(slide, 7)
    return slide


def slide_08_final_results(prs):
    slide = light_slide(prs)
    add_header(slide, "Validation", "Performance du modele retenu sur le test")

    top_stats = [
        ("0,841", "AUC"),
        ("79,4 %", "Rappel churn"),
        ("51,4 %", "Precision churn"),
        ("0,624", "F1-score"),
        ("74,6 %", "Accuracy"),
    ]
    x = Inches(0.35)
    box_w = Inches(1.82)
    for value, label in top_stats:
        add_stat_tile(slide, x, Inches(1.3), box_w, Inches(0.85), value, label)
        x = Emu(int(x) + int(box_w) + int(Inches(0.08)))

    add_textbox(
        slide,
        Inches(0.5),
        Inches(2.4),
        Inches(9),
        Inches(0.3),
        "Matrice de confusion (1 409 clients test)",
        size=13,
        bold=True,
        color=NAVY,
    )

    mat_left = Inches(1.6)
    mat_top = Inches(2.75)
    tbl = add_table(
        slide,
        mat_left,
        mat_top,
        Inches(6.8),
        Inches(1.65),
        ["", "Predit : Non-churn", "Predit : Churn"],
        [
            ["Reel : Non-churn", "754 (VN)", "281 (FP)"],
            ["Reel : Churn", "77 (FN)", "297 (VP)"],
        ],
        font_size=12.5,
    )
    # Recolorer les cellules selon exactitude (VN/VP corrects, FP/FN erreurs)
    coloring = {
        (1, 1): (GREEN_BG, GREEN_TEXT),
        (1, 2): (CORAL_BG, CORAL_TEXT),
        (2, 1): (CORAL_BG, CORAL_TEXT),
        (2, 2): (GREEN_BG, GREEN_TEXT),
    }
    for (r, c), (bg, fg) in coloring.items():
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = bg
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.color.rgb = fg
        run.font.bold = True

    add_textbox(
        slide,
        Inches(0.5),
        Inches(4.65),
        Inches(9),
        Inches(0.5),
        "297 clients partis detectes sur 374, au prix de 281 fausses alertes (77 departs manques).",
        size=12,
        color=DARK_TEXT,
    )
    add_footer(slide, 8)
    return slide


def slide_09_explainability(prs):
    slide = light_slide(prs)
    add_header(slide, "Explicabilite", "Comprendre chaque prediction (SHAP) — D-02")

    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.3),
        Inches(4.4),
        Inches(0.3),
        "Explicabilite par client",
        size=13,
        bold=True,
        color=NAVY,
    )
    add_bullets(
        slide,
        Inches(0.5),
        Inches(1.65),
        Inches(4.4),
        Inches(2.3),
        [
            "Le dashboard affiche, pour un client selectionne, les facteurs SHAP "
            "qui augmentent ou reduisent son risque de depart.",
            "Explicabilite locale (un client) et globale (l'ensemble du modele).",
        ],
        size=12.5,
    )

    add_textbox(
        slide,
        Inches(5.1),
        Inches(1.3),
        Inches(4.4),
        Inches(0.3),
        "Variables les plus influentes",
        size=13,
        bold=True,
        color=NAVY,
    )
    add_bullets(
        slide,
        Inches(5.1),
        Inches(1.65),
        Inches(4.4),
        Inches(2.3),
        [
            "Type de contrat (mensuel a risque, deux ans protecteur)",
            "Anciennete du client (tenure)",
            "MonthlyCharges et TotalCharges",
            "Service Internet fibre optique",
            "Absence de securite en ligne ou de support technique",
        ],
        size=12,
    )

    add_quote_box(
        slide,
        Inches(0.5),
        Inches(4.1),
        Inches(9),
        Inches(0.9),
        "Ce sont des associations apprises par le modele, pas des relations de "
        "cause a effet.",
    )
    add_footer(slide, 9)
    return slide


def slide_10_dashboard(prs):
    slide = light_slide(prs)
    add_header(slide, "Demonstration", "Explorer les predictions dans Streamlit")

    cards = [
        (
            "Priorisation & synthese",
            "Indicateurs temps reel, seuil interactif, liste de clients prioritaires",
        ),
        (
            "Explicabilite (SHAP)",
            "Facteurs de risque par client selectionne — competence D-02",
        ),
        (
            "Simulateur What-If",
            "Effet immediat d'un changement de contrat, tarif ou service offert",
        ),
        (
            "Monitoring & Data Drift",
            "Stabilite des donnees en production, rapport Evidently — competence D-06",
        ),
    ]
    card_w = Inches(4.35)
    card_h = Inches(1.35)
    gap = Inches(0.3)
    positions = [
        (Inches(0.5), Inches(1.35)),
        (Emu(int(Inches(0.5)) + int(card_w) + int(gap)), Inches(1.35)),
        (Inches(0.5), Emu(int(Inches(1.35)) + int(card_h) + int(gap))),
        (
            Emu(int(Inches(0.5)) + int(card_w) + int(gap)),
            Emu(int(Inches(1.35)) + int(card_h) + int(gap)),
        ),
    ]
    for (title, desc), (x, y) in zip(cards, positions):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_w, card_h)
        card.adjustments[0] = 0.07
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT_GRAY
        card.line.color.rgb = NAVY
        card.line.width = Pt(0.75)
        card.shadow.inherit = False
        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        p1 = tf.paragraphs[0]
        r1 = p1.add_run()
        r1.text = title
        r1.font.size = Pt(13)
        r1.font.bold = True
        r1.font.color.rgb = NAVY
        r1.font.name = "Calibri"
        p2 = tf.add_paragraph()
        p2.space_before = Pt(3)
        r2 = p2.add_run()
        r2.text = desc
        r2.font.size = Pt(11)
        r2.font.color.rgb = DARK_TEXT
        r2.font.name = "Calibri"

    add_textbox(
        slide,
        Inches(0.5),
        Inches(4.5),
        Inches(9),
        Inches(0.35),
        "streamlit run app.py  →  http://localhost:8501",
        size=12,
        bold=True,
        color=MID_GRAY,
    )
    add_footer(slide, 10)
    return slide


def slide_11_quality_security(prs):
    slide = light_slide(prs)
    add_header(slide, "Qualite logicielle", "Code teste, securise et conforme RGPD")

    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.3),
        Inches(4.4),
        Inches(0.3),
        "Securite & conformite RGPD",
        size=13,
        bold=True,
        color=NAVY,
    )
    add_bullets(
        slide,
        Inches(0.5),
        Inches(1.65),
        Inches(4.4),
        Inches(1.9),
        [
            "Anonymisation de customerID par HMAC-SHA256 salee",
            "Validation stricte des entrees via des schemas Pydantic",
            "Secrets isoles : .env non versionne, .env.example fourni",
        ],
        size=12.5,
    )

    add_textbox(
        slide,
        Inches(5.1),
        Inches(1.3),
        Inches(4.4),
        Inches(0.3),
        "Qualite & integration continue (D-04)",
        size=13,
        bold=True,
        color=NAVY,
    )
    add_bullets(
        slide,
        Inches(5.1),
        Inches(1.65),
        Inches(4.4),
        Inches(1.9),
        [
            "4 fichiers de tests : validation, pipeline, performance, securite",
            "Non-regression ML : Rappel ≥ 0,75, AUC ≥ 0,80",
            "CI/CD GitHub Actions : flake8 + black + pytest a chaque push",
        ],
        size=12.5,
    )

    labels = [
        "D-02 — Algorithme d'IA explicable",
        "D-04 — CI/CD automatisee",
        "D-06 — Monitoring en production",
    ]
    x = Inches(0.5)
    w = Inches(2.93)
    for label in labels:
        add_pill(
            slide,
            x,
            Inches(3.85),
            w,
            Inches(0.55),
            label,
            fill=NAVY,
            color=WHITE,
            size=10,
        )
        x = Emu(int(x) + int(w) + int(Inches(0.1)))
    add_footer(slide, 11)
    return slide


def slide_12_monitoring(prs):
    slide = light_slide(prs)
    add_header(
        slide, "Monitoring — D-06", "Surveiller la derive des donnees en production"
    )

    add_bullets(
        slide,
        Inches(0.5),
        Inches(1.4),
        Inches(9),
        Inches(1.9),
        [
            "Evidently AI compare la distribution des donnees de production a la "
            "distribution d'entrainement de reference.",
            "Un rapport interactif est genere (docs/data_drift_report.html) et "
            "accessible depuis l'onglet Monitoring du dashboard.",
            "Objectif : declencher une alerte, puis une investigation ou un "
            "reentrainement, si une derive significative est detectee.",
        ],
        size=13.5,
    )

    add_quote_box(
        slide,
        Inches(0.5),
        Inches(3.5),
        Inches(9),
        Inches(1.1),
        "python src/monitoring/drift.py — le rapport signale, variable par "
        "variable, les ecarts de distribution par rapport a la reference "
        "d'entrainement.",
    )
    add_footer(slide, 12)
    return slide


def slide_13_limits(prs):
    slide = light_slide(prs)
    add_header(slide, "Limites", "Ce que le modele ne montre pas")

    add_bullets(
        slide,
        Inches(0.5),
        Inches(1.35),
        Inches(4.4),
        Inches(3.0),
        [
            "Association ne signifie pas causalite",
            "Un seul decoupage train / test",
            "Dataset historique, population et periode limitees",
        ],
        size=13,
    )
    add_bullets(
        slide,
        Inches(5.1),
        Inches(1.35),
        Inches(4.4),
        Inches(3.0),
        [
            "Seuil de decision (0,5) non calibre sur un cout metier reel",
            "Aucune validation temporelle effectuee",
            "Une validation externe (nouvelles donnees) reste necessaire",
        ],
        size=13,
    )
    add_footer(slide, 13)
    return slide


def slide_14_next_steps(prs):
    slide = light_slide(prs)
    add_header(slide, "Suite", "Prochaines etapes recommandees")

    add_bullets(
        slide,
        Inches(0.5),
        Inches(1.35),
        Inches(9),
        Inches(3.2),
        [
            "Validation temporelle sur de nouvelles periodes",
            "Calibration des probabilites produites par le modele",
            "Seuil de decision defini avec les equipes metier, selon le cout d'une campagne",
            "Test A/B d'une campagne de retention reelle",
            "Monitoring continu en production, avec reentrainement declenche par la derive",
            "Comparaison avec des modeles de boosting (XGBoost, LightGBM)",
        ],
        size=14,
    )
    add_footer(slide, 14)
    return slide


def slide_15_conclusion(prs):
    slide = dark_slide(prs)
    add_header(slide, "Pour conclure", "Un outil d'aide a la priorisation", dark=True)

    quote = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.45), Inches(9), Inches(1.45)
    )
    quote.adjustments[0] = 0.1
    quote.fill.solid()
    quote.fill.fore_color.rgb = NAVY
    quote.line.fill.background()
    quote.shadow.inherit = False
    tf = quote.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.3)
    tf.margin_right = Inches(0.3)
    p1 = tf.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    r1 = p1.add_run()
    r1.text = "ChurnScope est un outil d'aide a la priorisation, pas une boite noire."
    r1.font.size = Pt(16)
    r1.font.bold = True
    r1.font.color.rgb = WHITE
    r1.font.name = "Calibri"
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(6)
    r2 = p2.add_run()
    r2.text = (
        "Le Random Forest optimise detecte environ 4 churns sur 5 (79,4 % de "
        "rappel), avec un controle humain systematique et un suivi continu."
    )
    r2.font.size = Pt(12.5)
    r2.font.color.rgb = ICE
    r2.font.name = "Calibri"

    add_textbox(
        slide,
        Inches(0.5),
        Inches(3.15),
        Inches(9),
        Inches(0.3),
        "Competences RNCP couvertes",
        size=13,
        bold=True,
        color=WHITE,
    )
    comp = [
        ("D-02", "Algorithme d'IA explicable et accessible"),
        ("D-04", "Pipeline CI/CD automatise"),
        ("D-06", "Monitoring de la performance en production"),
    ]
    x = Inches(0.5)
    w = Inches(2.93)
    for code, label in comp:
        chip = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(3.55), w, Inches(0.85)
        )
        chip.adjustments[0] = 0.1
        chip.fill.solid()
        chip.fill.fore_color.rgb = RGBColor(0x22, 0x3B, 0x5C)
        chip.line.color.rgb = GOLD
        chip.line.width = Pt(0.75)
        chip.shadow.inherit = False
        tf2 = chip.text_frame
        tf2.word_wrap = True
        tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf2.margin_left = Inches(0.15)
        tf2.margin_right = Inches(0.15)
        p1c = tf2.paragraphs[0]
        p1c.alignment = PP_ALIGN.CENTER
        r1c = p1c.add_run()
        r1c.text = code
        r1c.font.size = Pt(13)
        r1c.font.bold = True
        r1c.font.color.rgb = GOLD
        r1c.font.name = "Calibri"
        p2c = tf2.add_paragraph()
        p2c.alignment = PP_ALIGN.CENTER
        r2c = p2c.add_run()
        r2c.text = label
        r2c.font.size = Pt(9.5)
        r2c.font.color.rgb = ICE
        r2c.font.name = "Calibri"
        x = Emu(int(x) + int(w) + int(Inches(0.1)))

    add_textbox(
        slide,
        Inches(0.5),
        Inches(4.75),
        Inches(6),
        Inches(0.4),
        "Merci pour votre attention. Questions ?",
        size=14,
        bold=True,
        color=WHITE,
    )
    add_textbox(
        slide,
        SLIDE_W - Inches(3.3),
        Inches(4.8),
        Inches(2.8),
        Inches(0.3),
        "github.com/jouvence13/ChurnScope",
        size=10,
        color=ICE,
        align=PP_ALIGN.RIGHT,
    )
    return slide


# --------------------------------------------------------------------------- #
# Point d'entree
# --------------------------------------------------------------------------- #


def build_presentation() -> Path:
    prs = new_presentation()
    for slide_fn in (
        slide_01_title,
        slide_02_cadrage,
        slide_03_data,
        slide_04_pipeline,
        slide_05_protocol,
        slide_06_comparison,
        slide_07_optimization,
        slide_08_final_results,
        slide_09_explainability,
        slide_10_dashboard,
        slide_11_quality_security,
        slide_12_monitoring,
        slide_13_limits,
        slide_14_next_steps,
        slide_15_conclusion,
    ):
        slide_fn(prs)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT_PATH))
    return OUTPUT_PATH


def main():
    path = build_presentation()
    print(f"Presentation generee avec succes : {path}")
    print("Nombre de diapositives : 15")


if __name__ == "__main__":
    main()
